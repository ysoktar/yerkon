"""CLI end-to-end: validate-config, run, build-workbook, validate-outputs, run-all."""
import json
import textwrap

import openpyxl
from click.testing import CliRunner

from locbench3d.cli import main

SMALL_CONFIG = textwrap.dedent(
    """
    template:
      width_m: 10.0
      length_m: 10.0
      height_m: 4.0
      frame_duration_s: 0.001
      guard_duration_s: 0.0002
    variables:
      method: ["UWB_SS_TWR", "TDOA"]
      anchor_count: [5]
      tag_count: [1]
      update_rate_hz: [1.0]
    n_repeats: 5
    seed: 0
    range_bin_edges_m: [0, 5, 10, 20]
    hard_requirements:
      max_3d_p95_m: 5.0
      min_availability: 0.5
    """
)


def test_validate_config_command(tmp_path):
    config_path = tmp_path / "cfg.yaml"
    config_path.write_text(SMALL_CONFIG)
    runner = CliRunner()
    result = runner.invoke(main, ["validate-config", str(config_path)])
    assert result.exit_code == 0
    assert "scenario" in result.output.lower()


def test_build_workbook_command_reports_locked_file_cleanly(tmp_path, monkeypatch):
    """A workbook file locked by another program (e.g. open in Excel) must
    produce a clean, actionable CLI error and a nonzero exit, not a raw
    traceback pointing into openpyxl/zipfile internals."""
    import openpyxl.workbook.workbook as wb_module

    config_path = tmp_path / "cfg.yaml"
    config_path.write_text(SMALL_CONFIG)
    out_dir = tmp_path / "out"

    runner = CliRunner()
    run_result = runner.invoke(
        main, ["run", "--config", str(config_path), "--out", str(out_dir)]
    )
    assert run_result.exit_code == 0, run_result.output

    def _raise_permission_error(self, filename):
        raise PermissionError(13, "Permission denied")

    monkeypatch.setattr(wb_module.Workbook, "save", _raise_permission_error)

    result = runner.invoke(
        main,
        [
            "build-workbook",
            "--tables-dir", str(out_dir / "tables"),
            "--out", str(out_dir / "workbook.xlsx"),
        ],
    )
    assert result.exit_code == 1
    assert "ERROR" in result.output
    assert "excel" in result.output.lower() or "another program" in result.output.lower()
    assert "Traceback" not in result.output


def test_run_all_command_produces_workbook_and_tables(tmp_path):
    config_path = tmp_path / "cfg.yaml"
    config_path.write_text(SMALL_CONFIG)
    out_dir = tmp_path / "out"
    runner = CliRunner()
    result = runner.invoke(
        main, ["run-all", "--config", str(config_path), "--out", str(out_dir)]
    )
    assert result.exit_code == 0, result.output
    assert (out_dir / "tables" / "master_comparison.csv").exists()
    assert (out_dir / "workbook.xlsx").exists()
    assert (out_dir / "tables" / "manifest.json").exists()

    wb = openpyxl.load_workbook(str(out_dir / "workbook.xlsx"))
    assert "master_comparison" in wb.sheetnames
    assert "dashboard" in wb.sheetnames

    manifest = json.loads((out_dir / "tables" / "manifest.json").read_text())
    assert manifest["master_comparison"]["row_count"] > 0


def test_run_all_reports_validation_failures_with_nonzero_exit(tmp_path, monkeypatch):
    config_path = tmp_path / "cfg.yaml"
    config_path.write_text(SMALL_CONFIG)
    out_dir = tmp_path / "out"
    runner = CliRunner()
    # A first successful run to prove the happy path still exits 0.
    result = runner.invoke(
        main, ["run-all", "--config", str(config_path), "--out", str(out_dir)]
    )
    assert result.exit_code == 0
