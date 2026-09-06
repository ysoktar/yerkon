"""CLI wiring for the optional MATLAB UWB waveform import."""
import textwrap

import openpyxl
from click.testing import CliRunner

from locbench3d.cli import main

CONFIG = textwrap.dedent(
    """
    template:
      width_m: 10.0
      length_m: 10.0
      height_m: 4.0
      frame_duration_s: 0.001
      guard_duration_s: 0.0002
    variables:
      method: ["UWB_SS_TWR"]
      anchor_count: [5]
      tag_count: [1]
      update_rate_hz: [1.0]
    n_repeats: 5
    seed: 0
    range_bin_edges_m: [0, 5, 10, 20]
    """
)

MATLAB_CSV = (
    "true_range_m,estimated_range_m,error_m\n"
    "1,1.02,0.02\n"
    "5,4.9,-0.1\n"
)


def test_run_all_populates_matlab_sheet_when_csv_supplied(tmp_path):
    config_path = tmp_path / "cfg.yaml"
    config_path.write_text(CONFIG)
    matlab_csv = tmp_path / "matlab.csv"
    matlab_csv.write_text(MATLAB_CSV)
    out_dir = tmp_path / "out"

    runner = CliRunner()
    result = runner.invoke(
        main,
        [
            "run-all",
            "--config", str(config_path),
            "--out", str(out_dir),
            "--matlab-uwb-csv", str(matlab_csv),
        ],
    )
    assert result.exit_code == 0, result.output
    assert (out_dir / "tables" / "matlab_uwb_waveform.csv").exists()

    wb = openpyxl.load_workbook(str(out_dir / "workbook.xlsx"))
    ws = wb["matlab_uwb_waveform"]
    assert ws.max_row == 3  # header + 2 rows
    header = [c.value for c in ws[1]]
    assert "evidence_type" in header
    idx = header.index("evidence_type")
    assert ws.cell(row=2, column=idx + 1).value == "MATLAB_WAVEFORM"


def test_run_all_leaves_matlab_sheet_as_placeholder_without_csv(tmp_path):
    config_path = tmp_path / "cfg.yaml"
    config_path.write_text(CONFIG)
    out_dir = tmp_path / "out"

    runner = CliRunner()
    result = runner.invoke(
        main, ["run-all", "--config", str(config_path), "--out", str(out_dir)]
    )
    assert result.exit_code == 0, result.output

    wb = openpyxl.load_workbook(str(out_dir / "workbook.xlsx"))
    ws = wb["matlab_uwb_waveform"]
    assert ws.max_row == 1
    assert "no content" in str(ws.cell(row=1, column=1).value).lower()
